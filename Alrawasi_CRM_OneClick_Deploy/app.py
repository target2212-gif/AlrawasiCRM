APP_VERSION='7.5.2'
import urllib.request, urllib.error
import tempfile
from flask import Flask,request,redirect,url_for,render_template,session,flash,jsonify,abort,Response, send_file
import hashlib, sqlite3, os, urllib.parse, csv, io, json, re
from datetime import datetime
from openpyxl import load_workbook, Workbook
from functools import wraps
from urllib.parse import urljoin, quote_plus
import requests
from bs4 import BeautifulSoup

def check_password_hash(stored,pw):
 try:
  salt,hexd=stored.split('$',1); return hashlib.pbkdf2_hmac('sha256',pw.encode(),salt.encode(),200000).hex()==hexd
 except: return False
def generate_password_hash(pw):
 import secrets; salt=secrets.token_hex(16); return salt+'$'+hashlib.pbkdf2_hmac('sha256',pw.encode(),salt.encode(),200000).hex()

app=Flask(__name__); app.secret_key=os.environ.get('SECRET_KEY','change-this-secret')

@app.template_filter('clean_company_name')
def clean_company_name(value):
 value=str(value or '').strip()
 return re.sub(r'\s*#\d+\s*$','',value).strip()

BASE=os.path.dirname(__file__)
SEED_DB=os.path.join(BASE,'crm.db')
DATA_DIR=os.environ.get('DATA_DIR', BASE)
os.makedirs(DATA_DIR, exist_ok=True)
DB=os.path.join(DATA_DIR,'crm.db')
if DB != SEED_DB and not os.path.exists(DB):
 import shutil
 shutil.copy2(SEED_DB, DB)
STATUS={'new':'جديد','contacted':'تم التواصل','interested':'مهتم','quote_requested':'طلب عرض سعر','quote_sent':'تم إرسال عرض السعر','followup':'متابعة','contracted':'تم التعاقد','not_interested':'غير مهتم','unreachable':'لا يمكن التواصل','postponed':'مؤجل'}
STATUS_CLASS={k:'status-'+k.replace('_','-') for k in STATUS}
SPECIALTIES={'contracting':'خدمات المقاولين والشهادات','advertising':'الدعاية والإعلان','general':'خدمات عامة'}
DEFAULT_MESSAGES={
 'contracting':'السلام عليكم، معك {rep_name} من تلال الرواسي. نقدم خدمات تصنيف شركات المقاولات، شهادة المحتوى المحلي، وشهادات الأيزو. يسعدنا خدمة {company_name} وتزويدكم بالتفاصيل والعرض المناسب.',
 'advertising':'السلام عليكم، معك {rep_name} من تلال الرواسي للدعاية والإعلان. نقدم لوحات المحلات التجارية، الاستيكرات، البنرات، بوثات المعارض، الهوية والمطبوعات. يسعدنا خدمة {company_name} وتقديم عرض مناسب لاحتياجكم.',
 'general':'السلام عليكم، معك {rep_name} من تلال الرواسي. يسعدنا التعرف على احتياج {company_name} وتقديم خدماتنا والحلول المناسبة لكم.'}

def db():
 x=sqlite3.connect(DB, timeout=30)
 x.row_factory=sqlite3.Row
 x.execute('PRAGMA busy_timeout=30000')
 return x

def ensure_schema():
 # Gunicorn starts multiple workers at the same time. Serialize schema migration
 # so two workers cannot try to add the same SQLite column simultaneously.
 d=db()
 try:
  d.execute('BEGIN IMMEDIATE')
  cols={r['name'] for r in d.execute('PRAGMA table_info(users)').fetchall()}
  if 'specialty' not in cols: d.execute("ALTER TABLE users ADD COLUMN specialty TEXT DEFAULT 'general'")
  if 'whatsapp_template' not in cols: d.execute("ALTER TABLE users ADD COLUMN whatsapp_template TEXT")
  ccols={r['name'] for r in d.execute('PRAGMA table_info(companies)').fetchall()}
  if 'business_area' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN business_area TEXT DEFAULT 'contracting'")
  if 'muqawil_member_no' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN muqawil_member_no TEXT")
  if 'source_url' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN source_url TEXT")
  if 'address' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN address TEXT")
  if 'data_source' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN data_source TEXT")
  if 'imported_at' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN imported_at TEXT")
  if 'classification_status' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN classification_status TEXT")
  if 'classification_grade' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN classification_grade TEXT")
  if 'latitude' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN latitude REAL")
  if 'longitude' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN longitude REAL")
  if 'logo_url' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN logo_url TEXT")
  if 'google_place_id' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN google_place_id TEXT")
  if 'phone_source' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN phone_source TEXT")
  if 'phone_verified_at' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN phone_verified_at TEXT")
  if 'phone_match_name' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN phone_match_name TEXT")
  if 'phone_match_address' not in ccols: d.execute("ALTER TABLE companies ADD COLUMN phone_match_address TEXT")
  for col,typ in [('muqawil_key','TEXT'),('muqawil_internal_id','TEXT'),('muqawil_contractor_type','TEXT'),('company_size','TEXT'),('account_status','TEXT'),('training_hours','INTEGER'),('main_contractor_count','INTEGER'),('subcontractor_count','INTEGER'),('source_region_id','TEXT')]:
   if col not in ccols: d.execute(f"ALTER TABLE companies ADD COLUMN {col} {typ}")
  d.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_companies_muqawil_key ON companies(muqawil_key) WHERE muqawil_key IS NOT NULL")
  d.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_companies_crm_number ON companies(crm_number) WHERE crm_number IS NOT NULL")
  d.execute("UPDATE companies SET business_area='contracting' WHERE business_area IS NULL OR business_area=''")
  d.execute('CREATE INDEX IF NOT EXISTS idx_companies_assigned ON companies(assigned_user_id)')
  d.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_companies_muqawil_member ON companies(muqawil_member_no) WHERE muqawil_member_no IS NOT NULL')
  d.execute('CREATE INDEX IF NOT EXISTS idx_interactions_company ON interactions(company_id, contacted_at)')
  d.commit()
 except Exception:
  d.rollback()
  raise
 finally:
  d.close()
ensure_schema()

def refresh_muqawil_classification_index():
 """Populate Muqawil classification only; never touch CRM contact/status/follow-up data."""
 path=os.path.join(BASE,'muqawil_classification_index.json')
 if not os.path.exists(path): return
 try:
  with open(path,encoding='utf-8') as f: info=json.load(f)
  d=db()
  try:
   for source_id,v in info.items():
    if not source_id: continue
    d.execute("""UPDATE companies
                SET classification_status=?, classification_grade=?,
                    muqawil_member_no=COALESCE(NULLIF(muqawil_member_no,''),?)
                WHERE source_id=?""",
              (v.get('status') or 'غير مصنفة',v.get('grade') or '',v.get('member') or None,source_id))
   # Records from Muqawil that are not present in the index remain explicitly unknown.
   d.execute("""UPDATE companies SET classification_status='غير محدد'
                WHERE (data_source='Muqawil' OR source LIKE '%مقاول%' OR source_url LIKE '%muqawil.org%')
                  AND (classification_status IS NULL OR classification_status='')""")
   d.commit()
  finally: d.close()
 except Exception as e:
  print('classification migration warning:',e)
refresh_muqawil_classification_index()

def login_required(f):
 @wraps(f)
 def w(*a,**k):
  if 'uid' not in session:return redirect(url_for('login'))
  return f(*a,**k)
 return w

def admin_required(f):
 @wraps(f)
 def w(*a,**k):
  if 'uid' not in session:return redirect(url_for('login'))
  if session.get('role')!='admin': abort(403)
  return f(*a,**k)
 return w

def current_user(d): return d.execute('SELECT * FROM users WHERE id=?',(session['uid'],)).fetchone()
def specialty_scope(user, alias='c'):
 if user and (user['specialty'] or 'general')=='contracting': return f"{alias}.business_area='contracting'", []
 return '1=1', []
def scoped_company(d,cid):
 u=current_user(d); cond,args=specialty_scope(u,'c')
 return d.execute(f'SELECT c.*,u.name rep FROM companies c LEFT JOIN users u ON u.id=c.assigned_user_id WHERE c.id=? AND {cond}',[cid]+args).fetchone()
def can_add_companies(user):
 return True

def wa_text(user_row, company_name):
 specialty=(user_row['specialty'] or 'general') if user_row else 'general'; template=(user_row['whatsapp_template'] or '').strip() if user_row else ''
 if not template: template=DEFAULT_MESSAGES.get(specialty,DEFAULT_MESSAGES['general'])
 return template.replace('{{rep_name}}','{rep_name}').replace('{{company_name}}','{company_name}').format(rep_name=user_row['name'] if user_row else session.get('name',''),company_name=company_name or 'شركتكم')
def wa_link(phone,text):
 if not phone:return None
 p=''.join(ch for ch in phone if ch.isdigit())
 if p.startswith('05'): p='966'+p[1:]
 elif p.startswith('5') and len(p)==9: p='966'+p
 return 'https://wa.me/'+p+'?text='+urllib.parse.quote(text)

@app.context_processor
def ctx(): return {'status':STATUS,'status_class':STATUS_CLASS,'specialties':SPECIALTIES,'user':session.get('name')}
@app.route('/login',methods=['GET','POST'])
def login():
 if request.method=='POST':
  d=db(); u=d.execute('SELECT * FROM users WHERE email=? AND active=1',(request.form['email'],)).fetchone(); d.close()
  if u and check_password_hash(u['password_hash'],request.form['password']): session.clear(); session.update(uid=u['id'],name=u['name'],role=u['role']); return redirect(url_for('dashboard'))
  flash('بيانات الدخول غير صحيحة')
 return render_template('login.html')

MUQAWIL_BASE='https://muqawil.org'
MUQAWIL_LIST='https://muqawil.org/ar/contractors'

def _norm_name(v):
 return re.sub(r'\s+',' ',(v or '').strip()).casefold()

def _muqawil_cards(html):
 soup=BeautifulSoup(html,'html.parser')
 out=[]
 # Contractor detail links have /ar/contractors/<id>/<id>
 for a in soup.select('a[href*="/ar/contractors/"]'):
  href=(a.get('href') or '').strip()
  name=' '.join(a.stripped_strings).strip()
  if not name or not re.search(r'/ar/contractors/\d+/\d+', href): continue
  url=urljoin(MUQAWIL_BASE,href)
  if not any(x['url']==url for x in out):
   out.append({'name':name,'url':url})
 return out

def _muqawil_detail(url):
 r=requests.get(url,timeout=25,headers={'User-Agent':'AlrawasiCRM/1.0 (+business CRM sync)'})
 r.raise_for_status()
 soup=BeautifulSoup(r.text,'html.parser')
 txt=' '.join(soup.stripped_strings)
 def grab(label, stop_labels):
  i=txt.find(label)
  if i<0:return ''
  val=txt[i+len(label):].strip()
  cuts=[val.find(x) for x in stop_labels if val.find(x)>=0]
  return val[:min(cuts)].strip(' :-') if cuts else val[:180].strip(' :-')
 member=grab('رقم العضويه',['العضوية','عضو منذ','حجم المنشأة'])
 phone=grab('رقم جوال المنشأة',['البريد الإلكتروني','المدينة','المنطقه'])
 city=grab('المدينة',['المنطقه','عنوان'])
 region=grab('المنطقه',['عنوان','طلب تعاقد'])
 address=grab('عنوان',['طلب تعاقد','التراخيص','الأنشطة'])
 title=(soup.find('h1').get_text(' ',strip=True) if soup.find('h1') else '')
 classification_status='غير مصنفة'; classification_grade=''
 if 'غير مصنف' not in txt and 'مصنف' in txt:
  classification_status='مصنفة'
  gm=re.search(r'الدرجة\s+([^\s،()]+)',txt)
  classification_grade=gm.group(1) if gm else 'غير محددة'
 return {'name':title,'member':member,'phone':phone,'city':city,'region':region,'address':address,'url':url,
         'classification_status':classification_status,'classification_grade':classification_grade}

def sync_muqawil(max_pages=8):
 d=db(); added=existing=seen=pages_ok=0; errors=[]
 try:
  for page in range(1,max_pages+1):
   try:
    r=requests.get(MUQAWIL_LIST,params={'page':page},timeout=25,headers={'User-Agent':'Mozilla/5.0 (compatible; AlrawasiCRM/1.0)'})
    if r.status_code in (401,403):
     errors.append(f'منصة مقاول رفضت الاتصال الآلي (HTTP {r.status_code}). لم يتم تعديل أي عميل حالي.')
     break
    if r.status_code==429:
     errors.append('منصة مقاول حدّت عدد الطلبات مؤقتاً (HTTP 429). لم يتم تعديل أي عميل حالي.')
     break
    if r.status_code!=200:
     errors.append(f'صفحة {page}: HTTP {r.status_code}'); continue
    cards=_muqawil_cards(r.text); pages_ok+=1
    if not cards: continue
   except requests.Timeout:
    errors.append(f'صفحة {page}: انتهت مهلة الاتصال'); continue
   except requests.RequestException as e:
    errors.append(f'صفحة {page}: خطأ اتصال {type(e).__name__}'); continue
   except Exception as e:
    errors.append(f'صفحة {page}: {type(e).__name__}: {str(e)[:100]}'); continue
   for card in cards:
    seen+=1
    try:
     info=_muqawil_detail(card['url']); name=(info.get('name') or card.get('name') or '').strip()
     if not name: errors.append(f"{card.get('url','')}: تعذر قراءة اسم الشركة"); continue
     member=(info.get('member') or '').strip() or None; row=None
     if member: row=d.execute('SELECT id FROM companies WHERE muqawil_member_no=?',(member,)).fetchone()
     if not row:
      row=d.execute("SELECT id FROM companies WHERE LOWER(TRIM(COALESCE(name_ar,'')))=LOWER(TRIM(?)) OR LOWER(TRIM(COALESCE(name_en,'')))=LOWER(TRIM(?)) LIMIT 1",(name,name)).fetchone()
     if row:
      existing+=1; continue
     cols={x['name'] for x in d.execute('PRAGMA table_info(companies)').fetchall()}
     vals={'name_ar':name,'name_en':'','activity':'مقاولات','city':info.get('city',''),'phone':info.get('phone',''),'whatsapp':info.get('phone',''),'email':'','website':'','cr_number':'','status':'جديد','business_area':'contracting','muqawil_member_no':member,'source_url':info.get('url',card['url']),'address':info.get('address',''),'data_source':'Muqawil','imported_at':datetime.now().strftime('%Y-%m-%d %H:%M:%S'),'classification_status':info.get('classification_status','غير مصنفة'),'classification_grade':info.get('classification_grade','')}
     keys=[k for k in vals if k in cols]
     d.execute(f"INSERT INTO companies ({','.join(keys)}) VALUES ({','.join('?' for _ in keys)})",tuple(vals[k] for k in keys))
     d.commit(); added+=1
    except requests.Timeout:
     d.rollback(); errors.append(f"{card.get('name','شركة')}: انتهت مهلة صفحة الشركة")
    except requests.RequestException as e:
     d.rollback(); errors.append(f"{card.get('name','شركة')}: خطأ اتصال {type(e).__name__}")
    except sqlite3.IntegrityError:
     d.rollback(); existing+=1
    except Exception as e:
     d.rollback(); errors.append(f"{card.get('name','شركة')}: {type(e).__name__}: {str(e)[:100]}")
  return {'added':added,'existing':existing,'seen':seen,'pages_ok':pages_ok,'errors':errors,'error_count':len(errors)}
 finally: d.close()

@app.post('/admin/sync-muqawil')
@admin_required
def admin_sync_muqawil():
 flash('المزامنة المباشرة متوقفة بسبب HTTP 403. استخدم صفحة تحديث بيانات مقاول JSON.')
 return redirect(url_for('import_muqawil_json'))

@app.post('/internal/sync-muqawil')
def sync_muqawil_internal():
 token=request.headers.get('X-Sync-Token','')
 expected=os.environ.get('MUQAWIL_SYNC_TOKEN','')
 if not expected or not secrets.compare_digest(token,expected):
  return {'ok':False},403
 return {'ok':True,**sync_muqawil()}

@app.get('/health')
def health(): return jsonify({'status':'ok'})
@app.get('/logout')
def logout(): session.clear(); return redirect(url_for('login'))

@app.get('/')
@login_required
def dashboard():
 d=db(); is_admin=session.get('role')=='admin'; me=current_user(d); cond,args=specialty_scope(me,'companies'); scope=' WHERE '+cond
 total=d.execute('SELECT COUNT(*) n FROM companies'+scope,args).fetchone()['n']
 counts={k:d.execute('SELECT COUNT(*) n FROM companies'+scope+' AND marketing_status=?',args+[k]).fetchone()['n'] for k in STATUS}
 if is_admin:
  follow=d.execute("SELECT COUNT(DISTINCT company_id) n FROM interactions WHERE next_followup_at IS NOT NULL AND datetime(next_followup_at)<=datetime('now','localtime')").fetchone()['n']
  reps=d.execute("SELECT u.id,u.name,u.role,COUNT(c.id) n,SUM(CASE WHEN c.marketing_status='contacted' THEN 1 ELSE 0 END) contacted,SUM(CASE WHEN c.marketing_status='interested' THEN 1 ELSE 0 END) interested,SUM(CASE WHEN c.marketing_status='contracted' THEN 1 ELSE 0 END) contracted,(SELECT MAX(i.contacted_at) FROM interactions i WHERE i.user_id=u.id) last_contact FROM users u LEFT JOIN companies c ON c.assigned_user_id=u.id WHERE u.active=1 GROUP BY u.id ORDER BY n DESC").fetchall()
 else:
  follow=d.execute("SELECT COUNT(DISTINCT i.company_id) n FROM interactions i JOIN companies c ON c.id=i.company_id WHERE c.assigned_user_id=? AND i.next_followup_at IS NOT NULL AND datetime(i.next_followup_at)<=datetime('now','localtime')",(session['uid'],)).fetchone()['n']; reps=[]
 recent=d.execute(f"SELECT c.id,c.name_ar,c.name_en,i.channel,i.outcome,i.contacted_at,i.next_followup_at FROM interactions i JOIN companies c ON c.id=i.company_id WHERE {specialty_scope(me,'c')[0]} ORDER BY i.contacted_at DESC LIMIT 10").fetchall(); d.close()
 return render_template('dashboard.html',counts=counts,total=total,follow=follow,reps=reps,recent=recent,is_admin=is_admin)

@app.get('/companies')
@login_required
def companies():
 q=request.args.get('q','').strip(); city=request.args.get('city','').strip(); st=request.args.get('status','').strip(); rep=request.args.get('rep','').strip(); cls=request.args.get('classification','').strip(); page=max(1,int(request.args.get('page',1))); per=50; where=[]; args=[]
 d0=db(); me0=current_user(d0); spec_cond,spec_args=specialty_scope(me0,'c'); d0.close(); where.append(spec_cond); args += spec_args
 if session.get('role')=='admin' and rep:
  if rep=='unassigned': where.append('c.assigned_user_id IS NULL')
  else: where.append('c.assigned_user_id=?'); args.append(rep)
 if q: where.append('(c.name_ar LIKE ? OR c.name_en LIKE ? OR c.phone_mobile LIKE ? OR c.phone_landline LIKE ? OR c.email LIKE ?)'); args += [f'%{q}%']*5
 if city: where.append('c.city=?'); args.append(city)
 if st: where.append('c.marketing_status=?'); args.append(st)
 if cls=='classified': where.append("c.classification_status='مصنفة'")
 elif cls=='unclassified': where.append("c.classification_status='غير مصنفة'")
 w=(' WHERE '+' AND '.join(where)) if where else ''; d=db(); total=d.execute('SELECT COUNT(*) n FROM companies c'+w,args).fetchone()['n']
 sql="""SELECT c.*,u.name rep,(SELECT i.contacted_at FROM interactions i WHERE i.company_id=c.id ORDER BY i.contacted_at DESC LIMIT 1) last_contact,(SELECT i.next_followup_at FROM interactions i WHERE i.company_id=c.id AND i.next_followup_at IS NOT NULL ORDER BY i.contacted_at DESC LIMIT 1) next_followup FROM companies c LEFT JOIN users u ON u.id=c.assigned_user_id"""+w+' ORDER BY c.id LIMIT ? OFFSET ?'
 rows=[dict(r) for r in d.execute(sql,args+[per,(page-1)*per]).fetchall()]; cities=d.execute("SELECT DISTINCT city FROM companies WHERE city IS NOT NULL AND city<>'' ORDER BY city").fetchall(); me=d.execute('SELECT * FROM users WHERE id=?',(session['uid'],)).fetchone(); users=d.execute("SELECT id,name FROM users WHERE active=1 ORDER BY name").fetchall() if session.get('role')=='admin' else []; d.close()
 for r in rows: r['wa_link']=wa_link(r.get('whatsapp') or r.get('phone_mobile'),wa_text(me,r.get('name_ar') or r.get('name_en')))
 return render_template('companies.html',rows=rows,total=total,page=page,per=per,q=q,city=city,status_filter=st,rep_filter=rep,classification_filter=cls,cities=cities,users=users,is_admin=session.get('role')=='admin',can_add=can_add_companies(me),my_specialty=(me['specialty'] or 'general'))

@app.post('/companies/assign')
@admin_required
def assign_companies():
 ids=[int(x) for x in request.form.getlist('company_ids') if x.isdigit()]; uid=request.form.get('assigned_user_id')
 if not ids: flash('حدد عميلاً واحداً على الأقل'); return redirect(request.referrer or url_for('companies'))
 d=db(); valid=d.execute('SELECT id FROM users WHERE id=? AND active=1',(uid,)).fetchone() if uid else None
 if uid and not valid: d.close(); abort(400)
 ph=','.join('?'*len(ids)); d.execute(f'UPDATE companies SET assigned_user_id=?,updated_at=CURRENT_TIMESTAMP WHERE id IN ({ph})',[uid or None]+ids); d.commit(); d.close(); flash(f'تم توزيع {len(ids)} عميل بنجاح'); return redirect(request.referrer or url_for('companies'))

@app.get('/company/<int:cid>')
@login_required
def company(cid):
 d=db(); c=scoped_company(d,cid)
 if not c: d.close(); abort(404)
 its=d.execute('SELECT i.*,u.name rep FROM interactions i LEFT JOIN users u ON u.id=i.user_id WHERE company_id=? ORDER BY contacted_at DESC',(cid,)).fetchall(); users=d.execute('SELECT * FROM users WHERE active=1').fetchall() if session.get('role')=='admin' else []; me=d.execute('SELECT * FROM users WHERE id=?',(session['uid'],)).fetchone(); d.close(); link=wa_link(c['whatsapp'] or c['phone_mobile'],wa_text(me,c['name_ar'] or c['name_en']))
 return render_template('company.html',c=c,its=its,users=users,wa_link=link,is_admin=session.get('role')=='admin')

@app.post('/company/<int:cid>/update')
@login_required
def update_company(cid):
 d=db(); c=scoped_company(d,cid)
 if not c: d.close(); abort(404)
 if session.get('role')=='admin':
  assigned=request.form.get('assigned_user_id') or None
  d.execute('UPDATE companies SET marketing_status=?,assigned_user_id=?,priority=?,marketing_notes=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(request.form['marketing_status'],assigned,request.form.get('priority','normal'),request.form.get('marketing_notes',''),cid))
 else:
  d.execute('UPDATE companies SET marketing_status=?,priority=?,marketing_notes=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(request.form['marketing_status'],request.form.get('priority','normal'),request.form.get('marketing_notes',''),cid)); # حفظ كل تحديث لحالة العميل في سجل التواصل الحالي
 new_status=(request.form.get('marketing_status') or '').strip()
 if new_status:
  d.execute('INSERT INTO interactions(company_id,user_id,channel,outcome,notes) VALUES(?,?,?,?,?)',
            (cid,session['uid'],'تحديث حالة العميل',new_status,(request.form.get('marketing_notes') or '').strip()))
 d.commit(); d.close(); return redirect(url_for('company',cid=cid))

@app.post('/company/<int:cid>/interaction')
@login_required
def interaction(cid):
 d=db(); c=scoped_company(d,cid)
 if not c: d.close(); abort(404)
 d.execute('INSERT INTO interactions(company_id,user_id,channel,outcome,notes,next_followup_at) VALUES(?,?,?,?,?,?)',(cid,session['uid'],request.form['channel'],request.form['outcome'],request.form.get('notes',''),request.form.get('next_followup_at') or None)); d.execute('UPDATE companies SET marketing_status=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(request.form['outcome'],cid)); d.commit(); d.close(); return redirect(url_for('company',cid=cid))

@app.post('/company/<int:cid>/quick-followup')
@login_required
def quick_followup(cid):
 d=db(); c=scoped_company(d,cid)
 if not c: d.close(); abort(404)
 when=request.form.get('next_followup_at') or None; notes=request.form.get('notes','متابعة سريعة')
 d.execute("INSERT INTO interactions(company_id,user_id,channel,outcome,notes,next_followup_at) VALUES(?,?,?,?,?,?)",(cid,session['uid'],'متابعة','followup',notes,when)); d.execute("UPDATE companies SET marketing_status='followup',updated_at=CURRENT_TIMESTAMP WHERE id=?",(cid,)); d.commit(); d.close(); flash('تم تسجيل المتابعة السريعة'); return redirect(request.referrer or url_for('companies'))

@app.get('/users')
@admin_required
def users():
 d=db(); rows=d.execute('SELECT id,name,email,role,active,specialty,whatsapp_template FROM users ORDER BY id').fetchall(); d.close(); return render_template('users.html',rows=rows,default_messages=DEFAULT_MESSAGES)
@app.post('/users/add')
@admin_required
def add_user():
 specialty=request.form.get('specialty','general'); template=request.form.get('whatsapp_template','').strip() or DEFAULT_MESSAGES.get(specialty,DEFAULT_MESSAGES['general']); d=db(); d.execute('INSERT INTO users(name,email,password_hash,role,specialty,whatsapp_template) VALUES(?,?,?,?,?,?)',(request.form['name'],request.form['email'],generate_password_hash(request.form['password']),request.form['role'],specialty,template)); d.commit(); d.close(); return redirect(url_for('users'))
@app.post('/users/<int:uid>/update')
@admin_required
def update_user(uid):
 specialty=request.form.get('specialty','general'); template=request.form.get('whatsapp_template','').strip() or DEFAULT_MESSAGES.get(specialty,DEFAULT_MESSAGES['general']); d=db(); d.execute('UPDATE users SET specialty=?,whatsapp_template=? WHERE id=?',(specialty,template,uid)); d.commit(); d.close(); flash('تم تحديث اختصاص ورسالة المستخدم'); return redirect(url_for('users'))

@app.route('/companies/add',methods=['GET','POST'])
@login_required
def add_company():
 d=db(); me=current_user(d)
 if not can_add_companies(me): d.close(); abort(403)
 if request.method=='POST':
  area=request.form.get('business_area','advertising'); area=area if area in ('contracting','advertising') else 'advertising'
  vals=[request.form.get(k,'').strip() for k in ('name_ar','name_en','sector','subcategory','city','district','address','phone_mobile','phone_landline','whatsapp','email','website','cr_number')]
  d.execute("INSERT INTO companies(name_ar,name_en,sector,subcategory,city,district,address,phone_mobile,phone_landline,whatsapp,email,website,cr_number,source,marketing_status,business_area,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",vals+['manual','new',area]); d.commit(); d.close(); flash('تمت إضافة الشركة بنجاح'); return redirect(url_for('companies'))
 d.close(); return render_template('company_add.html')

@app.route('/companies/import',methods=['GET','POST'])
@login_required
def import_companies():
 d=db(); me=current_user(d)
 if not can_add_companies(me): d.close(); abort(403)
 if request.method=='POST':
  f=request.files.get('file'); area=request.form.get('business_area','advertising'); area=area if area in ('contracting','advertising') else 'advertising'
  if not f or not f.filename: d.close(); flash('اختر ملفاً'); return redirect(request.url)
  ext=os.path.splitext(f.filename.lower())[1]; records=[]
  aliases={'اسم الشركة':'name_ar','الاسم':'name_ar','name_ar':'name_ar','name':'name_ar','company':'name_ar','name_en':'name_en','الاسم الانجليزي':'name_en','القطاع':'sector','sector':'sector','النشاط':'sector','التصنيف':'subcategory','subcategory':'subcategory','المدينة':'city','city':'city','الحي':'district','district':'district','العنوان':'address','address':'address','الجوال':'phone_mobile','الهاتف المحمول':'phone_mobile','phone_mobile':'phone_mobile','mobile':'phone_mobile','الهاتف':'phone_landline','phone':'phone_landline','phone_landline':'phone_landline','واتساب':'whatsapp','whatsapp':'whatsapp','البريد':'email','البريد الإلكتروني':'email','email':'email','الموقع':'website','website':'website','السجل التجاري':'cr_number','cr_number':'cr_number'}
  try:
   if ext=='.csv': records=list(csv.DictReader(io.StringIO(f.read().decode('utf-8-sig'))))
   elif ext in ('.xlsx','.xlsm'):
    ws=load_workbook(f,read_only=True,data_only=True).active; vals=list(ws.iter_rows(values_only=True)); headers=[str(x or '').strip() for x in vals[0]]; records=[dict(zip(headers,row)) for row in vals[1:]]
   else: raise ValueError('صيغة غير مدعومة. استخدم CSV أو XLSX')
   n=0
   for row in records:
    x={}
    for k,v in row.items():
     key=aliases.get(str(k or '').strip().lower()) or aliases.get(str(k or '').strip())
     if key: x[key]=str(v or '').strip()
    if not (x.get('name_ar') or x.get('name_en')): continue
    vals=[x.get(k,'') for k in ('name_ar','name_en','sector','subcategory','city','district','address','phone_mobile','phone_landline','whatsapp','email','website','cr_number')]
    d.execute("INSERT INTO companies(name_ar,name_en,sector,subcategory,city,district,address,phone_mobile,phone_landline,whatsapp,email,website,cr_number,source,marketing_status,business_area,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",vals+['import','new',area]); n+=1
   d.commit(); d.close(); flash(f'تم استيراد {n} شركة بنجاح'); return redirect(url_for('companies'))
  except Exception as e: d.rollback(); d.close(); flash('تعذر الاستيراد: '+str(e)); return redirect(request.url)
 d.close(); return render_template('company_import.html')



def _extract_logo_url(x):
 """Return a company logo/image URL only when the Muqawil JSON actually supplies one."""
 for key in ('logoUrl','logo_url','logo','imageUrl','image_url','companyLogo','company_logo','avatar','photo'):
  v=x.get(key)
  if isinstance(v,str) and v.strip().lower().startswith(('http://','https://')):
   return v.strip()
 # Some exporters nest media/image fields.
 for parent in ('media','image','companyImage','company_image'):
  obj=x.get(parent)
  if isinstance(obj,dict):
   for key in ('url','src','logoUrl','imageUrl'):
    v=obj.get(key)
    if isinstance(v,str) and v.strip().lower().startswith(('http://','https://')):
     return v.strip()
 return ''

def _muqawil_catalog_records(payload):
 if isinstance(payload,list): return payload
 if isinstance(payload,dict) and isinstance(payload.get('companies'),list): return payload['companies']
 raise ValueError('ملف مقاول يجب أن يحتوي companies أو قائمة شركات')

def _muqawil_status(v):
 v=str(v or '').strip()
 if v in ('مصنف','مصنفة'): return 'مصنفة'
 if v in ('غير مصنف','غير مصنفة'): return 'غير مصنفة'
 return 'غير محدد'

def _next_crm_number(d):
 return (d.execute('SELECT COALESCE(MAX(crm_number),0)+1 FROM companies').fetchone()[0] or 1)

def _upsert_muqawil_v42(d,x,next_no):
 """Update Muqawil directory facts only. Never import phone/WhatsApp from collector files."""
 key=str(x.get('muqawilKey') or x.get('muqawil_key') or '').strip()
 iid=str(x.get('muqawilInternalId') or x.get('muqawil_internal_id') or '').strip()
 ctype=str(x.get('muqawilContractorType') or x.get('muqawil_contractor_type') or '').strip()
 profile=str(x.get('muqawilProfileUrl') or x.get('sourceUrl') or x.get('source_url') or '').strip()
 if not key and iid and ctype: key=f'{iid}:{ctype}'
 if not key:
  m=re.search(r'/ar/contractors/(\d+)/(\d+)',profile)
  if m: iid,ctype=m.group(1),m.group(2); key=f'{iid}:{ctype}'
 name=str(x.get('nameAr') or x.get('name_ar') or '').strip()
 if not key or not name: return 'rejected',next_no
 member=str(x.get('membershipNo') or x.get('muqawil_member_no') or '').strip() or None
 row=d.execute('SELECT id FROM companies WHERE muqawil_key=? LIMIT 1',(key,)).fetchone()
 if not row and profile: row=d.execute('SELECT id FROM companies WHERE muqawil_profile_url=? LIMIT 1',(profile,)).fetchone()
 vals=(name,str(x.get('city') or ''),str(x.get('address') or ''),_muqawil_status(x.get('classificationStatus') or x.get('classification_status')),str(x.get('classificationGrade') or x.get('classification_grade') or ''),member,str(x.get('logoUrl') or x.get('logo_url') or ''),profile,profile,key,iid,ctype,str(x.get('companySize') or x.get('company_size') or ''),str(x.get('accountStatus') or x.get('account_status') or ''),x.get('trainingHours'),x.get('mainContractorCount'),x.get('subcontractorCount'),str(x.get('sourceRegionId') or x.get('source_region_id') or ''))
 if row:
  d.execute('''UPDATE companies SET name_ar=?,city=?,address=?,classification_status=?,classification_grade=?,muqawil_member_no=?,logo_url=?,source_url=?,muqawil_profile_url=?,muqawil_key=?,muqawil_internal_id=?,muqawil_contractor_type=?,company_size=?,account_status=?,training_hours=?,main_contractor_count=?,subcontractor_count=?,source_region_id=?,source_verified=1,data_source='Muqawil',updated_at=CURRENT_TIMESTAMP WHERE id=?''',vals+(row['id'],))
  return 'existing',next_no
 d.execute('''INSERT INTO companies(source_id,name_ar,sector,city,address,source,source_url,source_verified,marketing_status,priority,classification_status,classification_grade,muqawil_member_no,business_area,data_source,imported_at,logo_url,muqawil_profile_url,muqawil_profile_status,crm_number,muqawil_key,muqawil_internal_id,muqawil_contractor_type,company_size,account_status,training_hours,main_contractor_count,subcontractor_count,source_region_id,phone_mobile,phone_landline,whatsapp,phone_source,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)''',
  (key,name,'مقاولات',str(x.get('city') or ''),str(x.get('address') or ''),'الهيئة السعودية للمقاولين (منصة مقاول)',profile,1,'new','normal',_muqawil_status(x.get('classificationStatus') or x.get('classification_status')),str(x.get('classificationGrade') or x.get('classification_grade') or ''),member,'contracting','Muqawil',datetime.now().strftime('%Y-%m-%d %H:%M:%S'),str(x.get('logoUrl') or x.get('logo_url') or ''),profile,'verified' if profile else 'missing',next_no,key,iid,ctype,str(x.get('companySize') or x.get('company_size') or ''),str(x.get('accountStatus') or x.get('account_status') or ''),x.get('trainingHours'),x.get('mainContractorCount'),x.get('subcontractorCount'),str(x.get('sourceRegionId') or x.get('source_region_id') or ''),'','','',''))
 return 'added',next_no+1

@app.route('/companies/import-muqawil',methods=['GET','POST'])
@admin_required
def import_muqawil_json():
 if request.method=='GET': return render_template('muqawil_import.html')
 f=request.files.get('file')
 if not f or not f.filename.lower().endswith('.json'):
  flash('اختر ملف JSON الخاص بمنصة مقاول.'); return redirect(request.url)
 d=None
 try:
  payload=json.load(f.stream); records=_muqawil_catalog_records(payload)
  d=db(); next_no=_next_crm_number(d); added=existing=rejected=0
  for x in records:
   if not isinstance(x,dict): rejected+=1; continue
   profile=str(x.get('muqawilProfileUrl') or x.get('sourceUrl') or '')
   if not (x.get('muqawilKey') or 'muqawil.org' in profile.lower()): rejected+=1; continue
   result,next_no=_upsert_muqawil_v42(d,x,next_no)
   if result=='added': added+=1
   elif result=='existing': existing+=1
   else: rejected+=1
  d.commit(); d.close(); d=None
  flash(f'تم تحديث دليل مقاول: جديد {added}، موجود/محدّث {existing}، مرفوض {rejected}. لم يتم استيراد أي أرقام هاتف من الملف.')
  return redirect(url_for('companies'))
 except Exception as e:
  if d: d.rollback(); d.close()
  flash('تعذر استيراد ملف مقاول: '+str(e)); return redirect(request.url)

@app.post('/admin/rebuild-muqawil-v42')
@admin_required
def rebuild_muqawil_v42():
 if request.form.get('confirm')!='REBUILD':
  flash('لم يتم التنفيذ. اكتب REBUILD للتأكيد.'); return redirect(url_for('import_muqawil_json'))
 path=os.path.join(BASE,'muqawil_v4_2_catalog.json')
 if not os.path.exists(path):
  flash('ملف دليل V4.2 المدمج غير موجود.'); return redirect(url_for('import_muqawil_json'))
 d=db()
 try:
  payload=json.load(open(path,encoding='utf-8')); records=_muqawil_catalog_records(payload)
  d.execute('DELETE FROM interactions'); d.execute('DELETE FROM companies')
  next_no=1; added=rejected=0
  for x in records:
   result,next_no=_upsert_muqawil_v42(d,x,next_no)
   if result=='added': added+=1
   elif result=='rejected': rejected+=1
  d.commit(); flash(f'تم إنشاء دليل مقاول V4.2: {added} شركة. المستخدمون والرسائل الجاهزة محفوظة، وأرقام الهواتف فارغة.')
 except Exception as e:
  d.rollback(); flash('تعذر إنشاء الدليل: '+str(e))
 finally: d.close()
 return redirect(url_for('companies'))


def _google_places_phone(company):
 key=(os.environ.get('GOOGLE_MAPS_API_KEY') or '').strip()
 if not key: return {'ok':False,'reason':'GOOGLE_MAPS_API_KEY غير مضبوط'}
 name=(company['name_ar'] or company['name_en'] or '').strip()
 city=(company['city'] or '').strip()
 address=(company['address'] or '').strip()
 query='، '.join(x for x in (name,address,city,'السعودية') if x)
 payload=json.dumps({'textQuery':query,'languageCode':'ar','regionCode':'SA','maxResultCount':3}).encode('utf-8')
 req=urllib.request.Request('https://places.googleapis.com/v1/places:searchText',data=payload,method='POST',
   headers={'Content-Type':'application/json','X-Goog-Api-Key':key,
            'X-Goog-FieldMask':'places.id,places.displayName,places.formattedAddress,places.nationalPhoneNumber,places.internationalPhoneNumber'})
 try:
  with urllib.request.urlopen(req,timeout=20) as r: data=json.load(r)
 except Exception as e:
  return {'ok':False,'reason':str(e)}
 places=data.get('places') or []
 if not places: return {'ok':False,'reason':'لا توجد نتيجة مطابقة'}
 # لا نختار نتيجة بعيدة: الاسم + المدينة/العنوان يجب أن يساندا المطابقة.
 def norm(s): return re.sub(r'[\W_]+','',str(s or '').lower(),flags=re.UNICODE)
 nn=norm(name); cc=norm(city)
 best=None
 for p in places:
  pn=norm((p.get('displayName') or {}).get('text',''))
  pa=norm(p.get('formattedAddress',''))
  name_ok=(nn and (nn in pn or pn in nn))
  city_ok=(not cc) or (cc in pa)
  phone=p.get('internationalPhoneNumber') or p.get('nationalPhoneNumber')
  if name_ok and city_ok and phone:
   best=p; break
 if not best: return {'ok':False,'reason':'المطابقة غير كافية أو لا يوجد هاتف'}
 return {'ok':True,'phone':best.get('internationalPhoneNumber') or best.get('nationalPhoneNumber'),
         'place_id':best.get('id'),'match_name':(best.get('displayName') or {}).get('text',''),
         'match_address':best.get('formattedAddress','')}

@app.post('/admin/google-phone/<int:cid>')
@admin_required
def google_phone_one(cid):
 d=db(); c=d.execute('SELECT * FROM companies WHERE id=?',(cid,)).fetchone()
 if not c: d.close(); abort(404)
 existing=(c['phone_mobile'] or c['phone_landline'] or '').strip()
 if existing:
  d.close(); flash('لم يتم تغيير الرقم: الشركة لديها رقم مسجل مسبقاً.'); return redirect(url_for('company',cid=cid))
 result=_google_places_phone(c)
 if result.get('ok'):
  d.execute("""UPDATE companies SET phone_mobile=?,phone_source='Google Maps / Places',
               google_place_id=?,phone_verified_at=CURRENT_TIMESTAMP,phone_match_name=?,phone_match_address=?
               WHERE id=?""",(result['phone'],result['place_id'],result['match_name'],result['match_address'],cid))
  d.commit(); flash('تمت إضافة رقم الهاتف من Google Maps بعد مطابقة الاسم والموقع.')
 else:
  flash('لم تتم إضافة رقم: '+result.get('reason','تعذر التحقق'))
 d.close(); return redirect(url_for('company',cid=cid))

@app.post('/admin/google-phone-batch')
@admin_required
def google_phone_batch():
 # دفعة صغيرة لحماية التكلفة والوقت؛ لا تستبدل أي رقم موجود.
 limit=min(max(int(request.form.get('limit') or 20),1),50)
 d=db()
 rows=d.execute("""SELECT * FROM companies
                   WHERE COALESCE(TRIM(phone_mobile),'')='' AND COALESCE(TRIM(phone_landline),'')=''
                   ORDER BY id LIMIT ?""",(limit,)).fetchall()
 added=skipped=0
 for c in rows:
  r=_google_places_phone(c)
  if r.get('ok'):
   d.execute("""UPDATE companies SET phone_mobile=?,phone_source='Google Maps / Places',
                google_place_id=?,phone_verified_at=CURRENT_TIMESTAMP,phone_match_name=?,phone_match_address=?
                WHERE id=?""",(r['phone'],r['place_id'],r['match_name'],r['match_address'],c['id']))
   added+=1
  else: skipped+=1
 d.commit(); d.close()
 flash(f'مطابقة Google Maps: أضيف {added} رقم، ولم تعتمد {skipped} نتيجة لعدم كفاية المطابقة/عدم وجود هاتف.')
 return redirect(url_for('companies'))

@app.get('/admin/backup-db')
@admin_required
def backup_db():
 """Download a transactionally consistent copy of the LIVE persistent SQLite database."""
 src=sqlite3.connect(DB,timeout=30)
 tmp=tempfile.NamedTemporaryFile(prefix='alrawasi_crm_',suffix='.db',delete=False)
 tmp.close()
 dst=sqlite3.connect(tmp.name)
 try:
  src.backup(dst)
 finally:
  dst.close(); src.close()
 return send_file(tmp.name,as_attachment=True,download_name='alrawasi_crm_live_backup.db',mimetype='application/octet-stream')

@app.get('/admin/export-full')
@admin_required
def export_full():
 """Export companies plus CRM history/users/followups to one Excel workbook."""
 d=db()
 wb=Workbook()
 # Remove default sheet and export every CRM table so history is not lost.
 wb.remove(wb.active)
 table_names=[r['name'] for r in d.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name").fetchall()]
 for table in table_names:
  ws=wb.create_sheet(title=table[:31])
  rows=d.execute(f'SELECT * FROM "{table}"').fetchall()
  cols=[r['name'] for r in d.execute(f'PRAGMA table_info("{table}")').fetchall()]
  ws.append(cols)
  for row in rows:
   ws.append([row[c] for c in cols])
  ws.freeze_panes='A2'
  ws.auto_filter.ref=ws.dimensions
 d.close()
 bio=io.BytesIO()
 wb.save(bio); bio.seek(0)
 return send_file(bio,as_attachment=True,download_name='alrawasi_crm_full_export.xlsx',
                  mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.get('/export')
@login_required
def export():
 d=db(); me=current_user(d); cond,args=specialty_scope(me,'companies'); rows=d.execute('SELECT * FROM companies WHERE '+cond+' ORDER BY id',args).fetchall(); out=io.StringIO(); w=csv.writer(out); w.writerow(rows[0].keys() if rows else ['id']); [w.writerow(list(r)) for r in rows]; d.close(); return Response('\ufeff'+out.getvalue(),mimetype='text/csv',headers={'Content-Disposition':'attachment; filename=companies.csv'})

@app.post('/company/<int:company_id>/delete')
@admin_required
def delete_company(company_id):
 d=db()
 c=d.execute('SELECT id,name_ar,name_en FROM companies WHERE id=?',(company_id,)).fetchone()
 if not c:
  d.close(); flash('الشركة غير موجودة.'); return redirect(url_for('companies'))
 try:
  d.execute('BEGIN IMMEDIATE')
  # Delete dependent rows from every table that has a company_id column.
  tables=[r['name'] for r in d.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()]
  for table in tables:
   if table=='companies': continue
   cols={r['name'] for r in d.execute(f'PRAGMA table_info("{table}")').fetchall()}
   if 'company_id' in cols:
    d.execute(f'DELETE FROM "{table}" WHERE company_id=?',(company_id,))
  d.execute('DELETE FROM companies WHERE id=?',(company_id,))
  d.commit()
  flash('تم حذف الشركة وجميع السجلات المرتبطة بها نهائياً.')
 except Exception:
  d.rollback(); raise
 finally:
  d.close()
 return redirect(url_for('companies'))

if __name__=='__main__': app.run(host='0.0.0.0',port=int(os.environ.get('PORT',5000)))
